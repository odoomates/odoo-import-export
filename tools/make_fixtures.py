"""Generate the workbook fixtures used by the unit tests (output is committed in tests/fixtures).

    python3 tools/make_fixtures.py      # needs openpyxl
"""
import datetime
import zipfile
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures"


def xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name", "Email", "Birthday", "Since", "Credit", "Active", "Tags"])
    ws.append(["Acme", "info@acme.test", datetime.date(1990, 4, 3), datetime.datetime(2024, 3, 1, 14, 30), 1234.5, True, "VIP, Gold"])
    ws.append(["Café Zoë", None, None, None, 0.1 + 0.2, False, None])
    ws.append([None, None, None, None, None, None, None])
    ws.append(["Long\ntext", "x@y.test", datetime.date(2000, 1, 31), None, -7, True, "Silver"])
    ws["H10"] = "far cell"
    other = wb.create_sheet("Second")
    other.append(["a", "b"])
    other.append([1, 2])
    wb.save(OUT / "sample.xlsx")


CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
<office:body><office:spreadsheet>
<table:table table:name="Sheet1">
<table:table-row>
  <table:table-cell office:value-type="string"><text:p>Product</text:p></table:table-cell>
  <table:table-cell office:value-type="string"><text:p>Price</text:p></table:table-cell>
  <table:table-cell office:value-type="string"><text:p>Date</text:p></table:table-cell>
  <table:table-cell office:value-type="string"><text:p>Sold</text:p></table:table-cell>
  <table:table-cell table:number-columns-repeated="1020"/>
</table:table-row>
<table:table-row>
  <table:table-cell office:value-type="string"><text:p>Desk</text:p></table:table-cell>
  <table:table-cell office:value-type="float" office:value="99.9"><text:p>99,90</text:p></table:table-cell>
  <table:table-cell office:value-type="date" office:date-value="2024-02-29"><text:p>29/02/2024</text:p></table:table-cell>
  <table:table-cell office:value-type="boolean" office:boolean-value="true"><text:p>TRUE</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="2">
  <table:table-cell office:value-type="string" table:number-columns-repeated="2"><text:p>same</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="1048570"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
</table:table>
</office:spreadsheet></office:body></office:document-content>
"""


def ods():
    with zipfile.ZipFile(OUT / "sample.ods", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", CONTENT)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx()
    ods()
